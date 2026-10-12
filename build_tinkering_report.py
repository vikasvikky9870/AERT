from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path(r'D:\AERT\Tinkering_Lab_Report_Monarch_Fit_AERT_Revised.docx')
MON=r'D:\MONARCH FIT'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
sec.header_distance=Inches(.3); sec.footer_distance=Inches(.32)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.2); normal.font.color.rgb=RGBColor(38,48,60)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for nm,size,col in [('Title',28,'12304A'),('Heading 1',19,'12304A'),('Heading 2',13,'007C83'),('Heading 3',11,'24566B')]:
 st=styles[nm]; st.font.name='Aptos Display' if nm!='Heading 3' else 'Aptos'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(col); st.paragraph_format.keep_with_next=True
styles['Heading 1'].paragraph_format.space_before=Pt(14); styles['Heading 1'].paragraph_format.space_after=Pt(7)
styles['Heading 2'].paragraph_format.space_before=Pt(10); styles['Heading 2'].paragraph_format.space_after=Pt(4)
hp=sec.header.paragraphs[0]; hp.text='TINKERING LAB  |  MONARCH FIT + AERT'; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs: r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(90,105,118)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=fp.add_run('Tinkering Lab 23BHL306  •  '); run.font.size=Pt(8); run.font.color.rgb=RGBColor(90,105,118)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell(cell,text,bold=False,color=None,size=8.5):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
 r=p.add_run(str(text)); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size)
 if color: r.font.color.rgb=RGBColor.from_string(color)
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(headers, rows, widths=None, size=8.3):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
 for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True,'FFFFFF',size); shade(t.rows[0].cells[i],'12304A')
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for i,val in enumerate(row):
   set_cell(cells[i],val,size=size)
   if ri%2==1: shade(cells[i],'F1F6F7')
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)
 return t
def p(text='',style=None): return doc.add_paragraph(text,style=style)
def bullets(items):
 for it in items: p(it,'List Bullet')
def numbered(items):
 for it in items: p(it,'List Number')
def h(text,l=1): doc.add_heading(text,level=l)
def caption(text):
 x=p(text); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in x.runs: r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(88,103,114)
def image(path,width=5.9):
 x=doc.add_picture(path,width=Inches(width)); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
def page(): doc.add_page_break()

# Cover
x=doc.add_paragraph(); x.paragraph_format.space_before=Pt(44); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=x.add_run('TINKERING LAB REPORT'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(27); r.font.color.rgb=RGBColor(18,48,74)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=x.add_run('Monarch Fit and Automated Emergency Response and Tracking'); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(0,124,131)
x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=x.add_run('Two applied technology prototypes for fitness guidance and emergency coordination'); r.italic=True; r.font.size=Pt(11)
doc.add_paragraph()
table(['COURSE','ACADEMIC DETAILS'],[['Tinkering Lab 23BHL306','3rd Year  |  5th Semester  |  CSE'],['Faculty Mentor','M. Divya'],['Module I','Monarch Fit — AI-assisted fitness training and gamification'],['Module II','AERT — emergency response and traffic intelligence prototype']], [1.45,5.5],9)
h('Student Team',2)
table(['Student Name','Roll Number'],[['BONDA AMERANDRA','24A51A05E3'],['BURAKA VIKAS','24A51A05E4'],['CHEVITI DILEEP KUMAR','24A51A05E5'],['CHINTADA REKHA RAGHAVI','24A51AO5E6'],['DEVADI PARNITHA','24A51A05E7']], [3.5,3.45],9)
p('Prepared with reference to the supplied project report, presentation, and current project repositories. Implementation status and limitations are stated separately from proposed future capabilities.').alignment=WD_ALIGN_PARAGRAPH.CENTER
page()
h('Contents')
for item in ['Abstract','1. Introduction and project overview','2. Tools, resources and evaluation approach','MODULE I — MONARCH FIT','3. Project summary and problem statement','4. Users, solution and workflow','5. Application architecture and technology','6. Evaluation plan, risks and future scope','MODULE II — AERT','7. Project summary and problem statement','8. Users, proposed solution and workflow','9. System architecture and technology','10. Prototype status and evaluation plan','11. Overall conclusion and learning outcomes','12. References and project materials']:
 p(item)
h('Abstract')
p('This report documents two student-developed software prototypes. Monarch Fit is an Android fitness application that combines workout planning, guided training, nutrition and activity tracking, and game-inspired progression. AERT (Automated Emergency Response and Tracking) is a geographic emergency-operations prototype that brings incident records, map-based situational awareness, resource views, and operator review into one interface. The projects address different problems but share a design goal: make complex tasks easier to understand through organized information and clear workflows.')
p('The report describes each project’s target users, problem, solution, workflow, technology, prototype evidence, evaluation plan, current limitations, and future work. The Monarch Fit section follows the current Android repository and the supplied presentation. The AERT section follows the current repository guide, which identifies the application as a training/demo system with seeded records and simulated telemetry. No user-study outcomes or live operational integrations are claimed here.')
h('Project overview')
table(['','MODULE I — MONARCH FIT','MODULE II — AERT'],[['Purpose','Support consistent fitness routines through guided training, tracking, and progression.','Help emergency operations staff review incidents, resources, hazards, and response options.'],['Primary users','Fitness app users, including beginners and enthusiasts.','Emergency operations staff in a training or demonstration setting.'],['Prototype form','Native Android application built with Kotlin and Jetpack Compose.','Web application with React/Vite frontend and FastAPI backend.'],['Current-state boundary','Fitness guidance and tracking prototype; device and feature behavior should be validated on target hardware.','Training/demo data; no connection to a real dispatch network, ambulance GPS feed, or hospital system.']], [1.1,2.9,3.0],8)

h('1. Introduction and project overview')
p('Digital tools can reduce friction when people need to make repeated decisions, but they must also present their limits clearly. Fitness applications need to connect plans, exercise guidance, activity records, and motivation so that users can follow a routine and understand their progress. Emergency coordination tools need to help staff assemble incident information, see resource context, and review options without implying that sample data or automated scores are verified facts.')
p('Monarch Fit explores the first challenge through a mobile training experience. AERT explores the second through a map-centered operations interface. The projects are not parts of one integrated system: they have separate users, workflows, data, and technology stacks. Their shared learning value lies in translating a real-world need into a prototype, making a workflow visible, and identifying what would need to be validated before wider use.')
h('2. Tools, resources and evaluation approach')
h('2.1 Software and project resources',2)
table(['Resource','Use in the projects'],[['Android Studio / Gradle / Kotlin toolchain','Build and maintain Monarch Fit as a native Android application.'],['Jetpack Compose and Material 3','Compose Monarch Fit screens and reusable interface components.'],['AERT React/Vite frontend','Develop the map-first emergency operations interface and workflow screens.'],['FastAPI, Python and SQLite','Provide local JSON endpoints, rule/scoring routines, and prototype persistence for AERT.'],['MapLibre and Esri-hosted raster tiles','Render AERT’s interactive map background and map layers; custom overlays remain separate from the basemap source.'],['Presentation and project reports','Communicate project objectives, workflows, prototype evidence, testing plans, and future directions.'],['Computers, smartphones and internet access','Development, review, mobile interaction checks, and project demonstrations.']], [2.05,4.95],8.4)
h('2.2 Evaluation method',2)
p('Evaluation should be task-based and should record what a participant can complete, where they hesitate, and whether the displayed information is understood. The team can recruit student participants, define short tasks for each module, observe completion and errors, and collect a short usability questionnaire. The materials reviewed for this revision did not provide completed participant counts, measured results, or survey summaries; therefore this report presents proposed tests and pass criteria rather than claiming test outcomes.')
table(['Evaluation dimension','Evidence to collect'],[['Task completion','Whether each participant completes the assigned task without facilitator intervention.'],['Clarity','Whether the participant can explain the next step and the meaning of the key status or recommendation.'],['Correctness','Whether records, totals, navigation state, and displayed outputs match the action taken.'],['Safety and honesty','Whether unavailable, simulated, approximate, or uncertain data is labelled clearly.'],['Performance','Loading, interaction response, readability, and layout on intended devices and screen sizes.']], [1.6,5.4],8.5)
page()

h('MODULE I — MONARCH FIT')
h('3. Project summary and problem statement')
h('3.1 Brief summary',2)
p('Monarch Fit is an Android-first fitness application designed to make training more engaging and easier to follow. Its core loop begins with a daily training decision, guides the user through a workout, records performance, and updates eligible progression. The current project README describes an Awaken setup flow, a Command hub with an explanation of the day’s training recommendation, an immersive Train session, a completion debrief, and an Evolve view for fitness attributes.')
p('The supplied presentation also describes exercise demonstrations, nutrition and activity tracking, quests, XP, streaks, and Hunter Ranks. The current repository identifies a Kotlin 2.0.21 and Jetpack Compose Android implementation, with a deterministic intelligence gateway, transparent “why” reasoning, physiological safety checks, and file-based persistence with idempotency protections. This report therefore uses those current repository details instead of carrying forward older Flutter implementation details found in the separate project report.')
h('3.2 Problem statement',2)
p('Many people find it difficult to maintain regular exercise because they are unsure what to do next, lack confidence in exercise technique, or lose motivation over time. When planning, training, nutrition, and progress are split across different tools, users must reconcile records themselves. A fitness interface can reduce that friction by presenting a clear plan, guided session controls, and progress feedback in a connected experience.')
h('3.3 Objectives',2)
bullets(['Present a clear daily workout recommendation and make the training action easy to find.','Guide users through active sets, repetitions, cadence, rest, and session completion.','Use a reusable 3D exercise avatar and movement clips where available.','Connect training history and progression to one consistent application state.','Use quests, XP, streaks, and fitness attributes to support motivation while preventing duplicate rewards.','Keep AI-assisted guidance explainable and bounded by safety rules.'])
h('4. Users, solution and workflow')
h('4.1 Intended users',2)
p('The intended users are people who want structured fitness routines and visible progress, including beginners who benefit from clear instructions and more experienced users who want a guided record of training. The application should not be presented as a medical diagnosis or a substitute for professional advice.')
h('4.2 Proposed solution',2)
p('Monarch Fit organizes the experience around five connected stages: Awaken establishes a user profile or training archetype; Command presents the day’s recommendation and its rationale; Train supports an active workout; Complete summarizes the finished session and eligible rewards; and Evolve visualizes progress across Strength, Endurance, Mobility, Recovery, Consistency, and Technique. The Android repository also includes models for workout, exercise, user profile, fitness statistics, and progression state.')
table(['Stage','User action','System response'],[['Awaken','Set up a fitness profile or select a training archetype.','Prepare relevant starting context for the fitness experience.'],['Command','Review the daily plan and its “why”.','Present the recommended protocol, streak, progression status, and entry to training.'],['Train','Record sets, repetitions, and session activity; follow cadence and rest cues.','Maintain active workout state and apply supported adaptations.'],['Complete','Finish the protocol.','Show a debrief, update eligible XP/progress, and avoid duplicate completion rewards.'],['Evolve','Review fitness attributes and progression.','Show progress across six fitness dimensions.']], [1.05,2.65,3.3],8.1)
h('4.3 Application workflow',2)
numbered(['Open the application and load saved profile and progression state.','Review the daily recommendation and the visible rationale.','Start or resume a workout and follow exercise instructions and available demonstrations.','Record sets, repetitions, cadence, and rest during the session.','Complete the session once; save the workout and update eligible progression.','Review the debrief and refreshed fitness attributes.'])
image(str(Path(MON)/'screen_command_updated.png'),3.7); caption('Figure 1. Monarch Fit Command screen from the project folder. Interface content illustrates the prototype experience.')
image(str(Path(MON)/'screen_train_live.png'),3.7); caption('Figure 2. Monarch Fit workout screen showing exercise guidance and session controls.')

h('5. Application architecture and technology')
p('The current project folder is a native Android codebase. The architecture separates feature screens, domain models, progression logic, and shared design components. The README describes a deterministic AI gateway pattern with decision validation and user-visible rationale; this is an explainable software decision layer, not evidence of a deployed general-purpose AI service.')
table(['Technology / component','Role in the current project'],[['Kotlin 2.0.21','Primary application language.'],['Jetpack Compose with Material 3','Android UI toolkit and component system.'],['Gradle 8.14 / Android Gradle Plugin 8.9.1','Build configuration described in the project README and app configuration.'],['Deterministic intelligence gateway','Coordinates local decision logic and exposes a rationale for recommendations.'],['Decision validator and safety rules','Check recommendations against defined physiological boundaries.'],['File serialization and idempotency protection','Persist application data and reduce duplicate state changes after repeated actions.'],['Compose Canvas visual components','Render progression and feedback visuals such as energy effects and fitness-stat charts.']], [2.25,4.75],8.1)
h('5.1 Prototype components',2)
bullets(['Awaken onboarding and fitness profile setup.','Command screen with today’s training directive, streak, level, and recommendation rationale.','Train screen with exercise, set/repetition targets, live cadence visualization, and rest timing.','Complete screen with session debrief and progression feedback.','Evolve screen with six-axis fitness attribute visualization.','Reusable Android UI and domain components for exercises, workouts, and progression.'])
p('The supplied project report contains older references to Flutter, GLB asset pipelines, 2,680 exercise catalogue entries, and Firebase. Those details were not carried into the description of the current build because the present Monarch Fit repository identifies a Kotlin/Compose Android implementation. Any features from the older report should be revalidated against the current source before reuse.')
h('6. Evaluation plan, risks and future scope')
h('6.1 Proposed test plan',2)
table(['ID','Task','Expected outcome / pass criterion'],[['M1','Open Command and explain the recommended session.','The participant identifies the next action and can describe the displayed rationale.'],['M2','Start a workout and progress through set, rep, cadence, and rest controls.','Session state advances correctly and remains understandable.'],['M3','Finish a workout and revisit progression.','History and eligible rewards update once; repeated completion does not duplicate rewards.'],['M4','Close and reopen during an active session.','Saved state follows the intended resume or discard behavior.'],['M5','Review Evolve attributes after a completed session.','Progress reflects supported saved activity and is legible.'],['M6','Trigger an unsafe or invalid recommendation input in a controlled test.','Safety validation rejects or bounds it and the reason is visible.']], [0.55,2.75,3.7],7.8)
h('6.2 Risks and safeguards',2)
bullets(['Fitness recommendations can be misunderstood as medical advice; use clear scope limits and avoid unsupported health claims.','Gamification can undermine trust if rewards are duplicated or detached from completed activity; base progression on validated state changes.','Recommendations must remain explainable; display the reason and do not imply medical measurement when the input is not available.','Workout persistence and interruption behavior should be checked on the target Android device.','Device, sensor, and external-service behavior should be described only after direct implementation and acceptance checks.'])
h('6.3 Future scope',2)
bullets(['Validate the current Android build and the full workout loop on physical devices.','Expand user-controlled personalization and make recommendation rationale easier to inspect.','Add more exercise demonstrations after movement accuracy and device playback are verified.','Improve accessibility, localization, and adjustable motion/visual settings.','Consider optional wearable or health-platform integrations with explicit permissions and data controls.'])
page()

h('MODULE II — AERT')
h('7. Project summary and problem statement')
h('7.1 Brief summary',2)
p('AERT (Automated Emergency Response and Tracking) is a geographic emergency-operations prototype focused on Srikakulam District, Andhra Pradesh. It brings an operations map, incident records, ambulance and hospital fixture records, road-event examples, workflow screens, and audit-style views into a browser application. The supplied presentation frames the concept around emergency call intake, triage support, ambulance assignment, route options, hospital matching, live tracking, and human-supervised authorization.')
p('The repository guide gives an essential current-state qualification: the application uses training/demo data and is not connected to a real emergency dispatch network. Some default-screen routes, traffic, weather, fleet movement, and facility information are hard-coded or seeded. The moving ambulance is animated, and a pre-alert action writes a local record rather than contacting a hospital. These distinctions are reflected throughout this report.')
h('7.2 Problem statement',2)
p('Emergency coordination can be slowed when incident details are gathered manually, locations are uncertain, road hazards are not visible to dispatchers, resource availability is unclear, or receiving facilities are contacted late. A centralized operations interface can help staff review information and coordinate a response, but only when the data sources, uncertainty, and authorization steps are explicit.')
h('7.3 Objectives',2)
bullets(['Bring incident, resource, road-event, and facility context into a geographic command interface.','Support structured intake and operator review of incident details.','Display response options in a way that can be reviewed and overridden by an authorized human.','Make the difference between sample/demo data and verified operational data visible.','Provide an auditable prototype workflow that can guide future integration and evaluation.'])
h('8. Users, proposed solution and workflow')
h('8.1 Intended users',2)
table(['User group','Need addressed by the prototype'],[['Emergency operations staff','A consolidated view of incident records, resources, map context, and workflow status.'],['Call takers / incident reviewers','A structured way to enter notes, confirm details, and identify missing information.'],['Dispatch decision-makers','A review surface for response options with human authorization.'],['Project evaluators / trainees','A controlled demonstration of how an emergency-operations workflow could be organized.']], [2.0,5.0],8.3)
h('8.2 Proposed solution',2)
p('AERT proposes a coordinated workflow that starts with incident information, makes location and context visible on a map, presents resource and route considerations, and keeps a dispatcher in the authorization loop. The platform’s screens include command/overwatch, call intake, incident views, fleet and hospital resources, geospatial intelligence, and audit or reporting surfaces. The slide deck’s multi-agent and quantum-inspired descriptions are treated as presentation-level concepts; the current repository guide describes deterministic rules, agent-shaped software components, and a classical QAOA-style numerical simulation rather than connected autonomous agents or quantum hardware.')
h('8.3 Operator-assisted workflow',2)
numbered(['Enter call notes or a transcript into the prototype’s intake workflow.','Review any rule-generated candidate clues and follow-up questions; use manual review if analysis is unavailable.','Confirm incident type, priority, casualty count, hazard, and location with the information available.','Choose or enter a map point and save an incident record.','Review map context, sample resources, and displayed response options.','Require an authorized operator to approve, modify, or reject any proposed response.','Record the outcome for training and review; no real dispatch or external hospital message is sent by this prototype.'])
image(r'D:\AERT\frontend\02_e2e_operator_workstation_1440.png',6.45); caption('Figure 3. AERT operations interface capture. This is a prototype view using seeded/demo records, not a live emergency control-room feed.')
image(r'D:\AERT\frontend\05_ai_recommendation.png',6.45); caption('Figure 4. AERT response-option review capture. Displayed route, unit, facility, confidence, and telemetry values are illustrative and require validation before operational use.')
h('9. System architecture and technology')
p('The current AERT repository contains a React/Vite browser frontend and a FastAPI backend with local SQLite persistence. MapLibre renders the interactive map; the inspected raster tile URLs are Esri services, while AERT’s incident and resource overlays have separate data provenance. Three.js is used for an optional procedural 3D visualization. Python modules provide rule-based analysis and numerical experiments. These components form a local prototype rather than a connected district-wide dispatch service.')
table(['Layer','Technology / current behavior'],[['Browser interface','React 19 with Vite; map and operations workflow components.'],['Map display','MapLibre GL JS with Esri-hosted raster tile services; custom pins, routes, and overlays are not validated by the basemap provider.'],['Optional 3D view','Three.js procedural visualization; not a surveyed or scale-accurate navigation model.'],['API','FastAPI with Pydantic validation for JSON endpoints.'],['Local persistence','SQLite database with seeded example records. Not a managed, replicated production datastore.'],['Decision support','Deterministic rules, scoring, and distance/time approximations; not a clinically validated triage or calibrated ETA service.'],['Optimization experiment','QUBO-like selection and QAOA-style statevector simulation computed classically with NumPy; no quantum device is connected.']], [1.65,5.35],8)
h('9.1 Data and integration boundary',2)
p('The project guide states that no connection is configured to 108/112 telephony, an ambulance GPS/AVL feed, a hospital information system, live traffic authority feeds, flood sensors, weather services, or a validated emergency road graph. The default interface may use seeded records or scripted scenarios. The alternate operational shell provides a more cautious manual intake and review path, but it is not currently the default experience. A map background or animated marker does not verify the incident pin, unit location, route, ETA, or facility capacity.')
h('10. Prototype status and evaluation plan')
h('10.1 Implemented prototype capabilities described by the repository guide',2)
bullets(['Map-centered interface with incident, ambulance, hospital, and road-event views.','Local incident records and API endpoints for workflow and analysis features.','Seeded sample data and scenario-based demonstration flows.','Operator-facing screens for incident review, resource views, and audit-style information.','Rule-based analysis and a classical numerical optimization experiment.'])
h('10.2 Test plan',2)
table(['ID','Task','Expected outcome / pass criterion'],[['A1','Create an incident with incomplete location details.','The UI exposes missing/uncertain information and requires operator confirmation before saving.'],['A2','Enter notes with negation or ambiguous wording.','Rule-generated clues are treated as candidates; the reviewer can correct them.'],['A3','Review a seeded resource and route recommendation.','The screen clearly labels sample data and does not imply verified live availability.'],['A4','Use a scenario or tracking animation.','The participant understands the action is simulated and not connected to a real vehicle.'],['A5','Trigger a hospital pre-alert action.','The interface identifies a local prototype record and does not imply the hospital was contacted.'],['A6','Reject or modify a displayed response option.','Human review remains possible and the action is recorded in the prototype workflow.'],['A7','Open the interface at desktop, tablet, and mobile widths.','Core information remains legible and controls do not overlap or become inaccessible.']], [0.55,2.75,3.7],7.7)
h('10.3 Safety, privacy and operational risks',2)
bullets(['AERT is not a replacement for approved emergency call-taking protocols or authorized dispatch systems.','Keyword matching can miss context, negation, multilingual phrasing, and speech-recognition errors; operators must confirm facts.','Straight-line distance and hard-coded routes do not provide validated road-network ETAs.','Sample fleet and hospital data must not be used to make real clinical or dispatch decisions.','Production use would require secure identity and access controls, encrypted data handling, audit retention, availability planning, and agreements for authoritative external data.','Emergency workflow decisions need operational and clinical governance before any automated recommendation is used in practice.'])
h('10.4 Future scope',2)
bullets(['Make the safer operator-reviewed intake the default experience and label all scenario/demo states consistently.','Integrate authoritative location, road, traffic, weather, fleet, and facility data only through approved agreements and with freshness indicators.','Evaluate intake and decision support against representative local language and historical cases under expert review.','Compare optimization experiments with transparent classical baselines using realistic, documented data.','Add authentication, role permissions, secure storage, retention controls, monitoring, backup, and disaster recovery before production consideration.','Conduct tabletop exercises with emergency operations stakeholders and measure task completion, error detection, and decision traceability.'])
h('11. Overall conclusion and learning outcomes')
p('Monarch Fit and AERT demonstrate two distinct uses of software to support human decision-making. Monarch Fit connects a daily plan, training session, saved activity, and progression loop in a native mobile experience. AERT organizes incident and resource context in a map-based emergency-operations prototype. Their value as Tinkering Lab projects is in making the intended workflows concrete and revealing the engineering and validation work needed to make those workflows dependable.')
p('The review also reinforced the importance of documenting current implementation accurately. Monarch Fit’s current project folder identifies a Kotlin/Compose application, while an older supplied report describes a Flutter version; the revised technical summary follows the current project folder. AERT’s screenshots and advanced slide descriptions are compelling demonstrations, but the repository guide clarifies that the operating data and integrations remain simulated or local. Clear labels, human review, measured evaluation, and honest limits are necessary parts of a responsible prototype.')
h('11.1 Learning outcomes',2)
bullets(['Translate a user problem into objectives, screens, workflow stages, and evaluation tasks.','Separate architecture and implemented features from concepts proposed for future work.','Use shared state and validation to keep application outputs consistent and trustworthy.','Recognize that visual realism does not establish data accuracy, live connectivity, or operational readiness.','Plan tests around usability, correctness, edge cases, and safety boundaries.'])
h('12. References and project materials')
for ref in ['Tinkering_Lab_Report_batch_2.docx, supplied report structure, team roster, and course details.','BATCH 3 TINKERING.pptx, supplied project framing, objectives, proposed workflows, and presentation claims.','Monarch_Fit_Full_Project_Report.docx, supplied project background, workflow concepts, and quality considerations. Older Flutter-specific implementation details were not treated as current because the project folder documents a newer Android build.','MONARCH FIT project folder: README.md, build.gradle.kts, app/build.gradle.kts, and app/src/main/java/com/monarchfit/ application sources.','AERT project folder: docs/AERT_PLATFORM_GUIDE.md and docs/AERT_BEGINNER_GUIDE.md; frontend and backend source code.','Prototype screenshots: MONARCH FIT/screen_command_updated.png; MONARCH FIT/screen_train_live.png; AERT/frontend/02_e2e_operator_workstation_1440.png; AERT/frontend/05_ai_recommendation.png.']:
 p(ref,'List Bullet')
p('Report date: 12 October 2026.')
for t in doc.tables:
 for row in t.rows:
  trPr=row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
doc.core_properties.title='Tinkering Lab Report — Monarch Fit and AERT'
doc.core_properties.subject='Project report for Monarch Fit and Automated Emergency Response and Tracking'
doc.core_properties.author='Tinkering Lab Project Team'
doc.save(OUT)
print(OUT)
